# -*- coding: utf-8 -*-
import docx
d = docx.Document(r"D:\黑客松\赛题 1：数据共情者-客服工作台示说明.docx")
for i, p in enumerate(d.paragraphs):
    if p.text.strip():
        print(f"[{p.style.name}] {p.text}")
for t_i, table in enumerate(d.tables):
    print(f"--- TABLE {t_i} ---")
    for row in table.rows:
        print(" | ".join(c.text.strip() for c in row.cells))
