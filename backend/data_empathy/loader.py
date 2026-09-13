"""Read-only Excel loader. Never writes the source workbook."""

from __future__ import annotations

import hashlib
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .codec import IdCoercionError, as_bool, as_int, as_number, as_string, as_string_id
from .mapping import (
    BUSINESS_SHEETS,
    CHAT_FIELDS,
    CHAT_SHEET,
    ORDER_FIELDS,
    ORDER_SHEET,
    WORKORDER_COMMON_FIELDS,
    WORKORDER_SHEETS,
)
from .timeutil import TimeParseError, parse_shanghai, time_annotation

XLSX_NAME = "赛题 1：数据共情者-业务数据.xlsx"


def repo_root() -> Path:
    here = Path(__file__).resolve()
    for parent in here.parents:
        if (parent / "AGENTS.md").exists() and (parent / "方案文档").is_dir():
            return parent
    raise FileNotFoundError("cannot locate repo root from data_empathy package")


def default_xlsx_path() -> Path:
    root = repo_root()
    path = root / XLSX_NAME
    if not path.is_file():
        raise FileNotFoundError(f"Excel not found: {path}")
    return path


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class ChatRecord:
    session_id: str
    message_no: int
    message_id: str
    message_time: datetime
    message_time_raw: str
    role: str
    buyer_nick: str
    sender: str | None
    shop: str | None
    scene_major: str | None
    scene_minor: str | None
    is_target_buyer_message: bool
    message_text: str
    content_type: str | None
    chat_content: str | None
    category: str | None
    image_path: str | None
    linked_order_id: str | None
    linked_workorder_id: str | None
    source_sheet: str
    source_row: int


@dataclass
class OrderRecord:
    order_id: str
    session_id: str
    buyer_nick: str
    shop: str | None
    sku_code: str | None
    product_name: str | None
    quantity: int | None
    unit_price: int | float | None
    paid_amount: int | float | None
    order_status_raw: str | None
    placed_time: datetime
    placed_time_raw: str
    paid_time: datetime | None
    paid_time_raw: str | None
    paid_time_note: str | None
    shipped_time: datetime | None
    shipped_time_raw: str | None
    courier: str | None
    tracking_no: str | None
    receiver_province: str | None
    receiver_city: str | None
    gift: str | None
    buyer_note: str | None
    source_sheet: str
    source_row: int


@dataclass
class WorkorderRecord:
    workorder_id: str
    session_id: str
    order_id: str | None
    buyer_nick: str | None
    shop: str | None
    kind: str
    source_sheet: str
    status_raw: str | None
    handler: str | None
    create_time: datetime
    create_time_raw: str
    complete_time: datetime | None
    complete_time_raw: str | None
    extra: dict[str, Any]
    source_row: int


@dataclass
class SheetLoad:
    name: str
    headers: list[str]
    row_count: int
    canonical_fields: list[str]
    parse_errors: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class CanonicalBundle:
    source_path: str
    source_sha256: str
    sheets: dict[str, SheetLoad]
    chats: list[ChatRecord]
    orders: list[OrderRecord]
    workorders: list[WorkorderRecord]
    log: list[str]


class DataStore:
    def __init__(self, bundle: CanonicalBundle):
        self.bundle = bundle
        chats_by_session: dict[str, list[ChatRecord]] = defaultdict(list)
        for row in bundle.chats:
            chats_by_session[row.session_id].append(row)
        for rows in chats_by_session.values():
            rows.sort(key=lambda r: (r.message_no, r.message_time, r.message_id))
        self.chats_by_session = dict(chats_by_session)
        self.chat_by_id = {r.message_id: r for r in bundle.chats}
        self.orders_by_session = {r.session_id: r for r in bundle.orders}
        self.orders_by_id = {r.order_id: r for r in bundle.orders}
        self.workorders_by_session = {r.session_id: r for r in bundle.workorders}
        self.workorders_by_id = {r.workorder_id: r for r in bundle.workorders}

    def require_session(self, session_id: str) -> list[ChatRecord]:
        rows = self.chats_by_session.get(session_id)
        if not rows:
            raise KeyError(f"unknown session_id: {session_id}")
        return rows

    def message(self, session_id: str, message_no: int) -> ChatRecord:
        for row in self.require_session(session_id):
            if row.message_no == message_no:
                return row
        raise KeyError(f"unknown message {session_id}#{message_no}")


def _header_index(headers: Iterable[Any]) -> dict[str, int]:
    index = {}
    for i, name in enumerate(headers):
        if name is None or str(name).strip() == "":
            continue
        key = str(name).strip()
        if key in index:
            raise ValueError(f"duplicate header {key!r}")
        index[key] = i
    return index


def _cell(row: tuple[Any, ...], index: dict[str, int], source: str):
    if source not in index:
        raise KeyError(f"missing column {source!r}")
    pos = index[source]
    if pos >= len(row):
        return None
    return row[pos]


def _convert(kind: str, value: Any, *, field: str):
    if kind == "string_id":
        return as_string_id(value, field=field)
    if kind == "string":
        return as_string(value)
    if kind == "int":
        return as_int(value, field=field)
    if kind == "number":
        return as_number(value, field=field)
    if kind == "bool":
        return as_bool(value, field=field)
    if kind == "time":
        return parse_shanghai(value, field=field)
    raise ValueError(f"unknown type {kind}")


def load_workbook_readonly(xlsx_path: str | Path | None = None) -> CanonicalBundle:
    path = Path(xlsx_path) if xlsx_path else default_xlsx_path()
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    digest = sha256_file(path)
    log = [
        f"open_readonly {path}",
        f"source_sha256 {digest}",
        "engine=openpyxl data_only=True read_only=True write=never",
    ]
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        present = list(wb.sheetnames)
        log.append(f"workbook_sheets {present}")
        missing = [name for name in BUSINESS_SHEETS if name not in wb.sheetnames]
        if missing:
            raise KeyError(f"missing business sheets: {missing}")
        sheets: dict[str, SheetLoad] = {}
        chats = _load_chat(wb[CHAT_SHEET], sheets, log)
        orders = _load_orders(wb[ORDER_SHEET], sheets, log)
        workorders: list[WorkorderRecord] = []
        for sheet_name in WORKORDER_SHEETS:
            workorders.extend(_load_workorders(wb[sheet_name], sheet_name, sheets, log))
    finally:
        wb.close()
        log.append("workbook_closed_without_save")
    return CanonicalBundle(
        source_path=str(path),
        source_sha256=digest,
        sheets=sheets,
        chats=chats,
        orders=orders,
        workorders=workorders,
        log=log,
    )


def _load_chat(ws, sheets: dict[str, SheetLoad], log: list[str]) -> list[ChatRecord]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("聊天记录 is empty")
    headers = [None if h is None else str(h).strip() for h in rows[0]]
    index = _header_index(headers)
    required = [src for src, _, _ in CHAT_FIELDS]
    missing = [name for name in required if name not in index]
    if missing:
        raise KeyError(f"聊天记录 missing columns: {missing}")
    records: list[ChatRecord] = []
    errors: list[dict[str, Any]] = []
    for excel_row, raw in enumerate(rows[1:], start=2):
        if _row_empty(raw):
            continue
        try:
            session_id = as_string_id(_cell(raw, index, "会话ID"), field="session_id")
            message_id = as_string_id(_cell(raw, index, "message_id"), field="message_id")
            message_no = as_int(_cell(raw, index, "消息序号"), field="message_no")
            if not session_id or not message_id or message_no is None:
                raise IdCoercionError(f"聊天记录 row {excel_row}: missing primary key")
            time_raw = _cell(raw, index, "发送时间")
            records.append(
                ChatRecord(
                    session_id=session_id,
                    message_no=message_no,
                    message_id=message_id,
                    message_time=parse_shanghai(time_raw, field="message_time"),
                    message_time_raw=as_string(time_raw) or "",
                    role=as_string(_cell(raw, index, "角色")) or "",
                    buyer_nick=as_string_id(_cell(raw, index, "买家昵称"), field="buyer_nick") or "",
                    sender=as_string(_cell(raw, index, "发送方")),
                    shop=as_string(_cell(raw, index, "店铺")),
                    scene_major=as_string(_cell(raw, index, "scene_major")),
                    scene_minor=as_string(_cell(raw, index, "scene_minor")),
                    is_target_buyer_message=bool(
                        as_bool(_cell(raw, index, "is_target_buyer_message"), field="is_target_buyer_message")
                    ),
                    message_text=as_string(_cell(raw, index, "message_text")) or "",
                    content_type=as_string(_cell(raw, index, "内容类型")),
                    chat_content=as_string(_cell(raw, index, "chat_content")),
                    category=as_string(_cell(raw, index, "category")),
                    image_path=as_string(_cell(raw, index, "image_path")),
                    linked_order_id=as_string_id(_cell(raw, index, "关联订单号"), field="linked_order_id"),
                    linked_workorder_id=as_string_id(_cell(raw, index, "关联工单号"), field="linked_workorder_id"),
                    source_sheet=CHAT_SHEET,
                    source_row=excel_row,
                )
            )
        except (IdCoercionError, TimeParseError, KeyError, ValueError) as exc:
            errors.append({"sheet": CHAT_SHEET, "row": excel_row, "error": str(exc)})
    sheets[CHAT_SHEET] = SheetLoad(
        name=CHAT_SHEET,
        headers=[h for h in headers if h],
        row_count=len(records),
        canonical_fields=[dst for _, dst, _ in CHAT_FIELDS],
        parse_errors=errors,
    )
    log.append(f"loaded {CHAT_SHEET} rows={len(records)} errors={len(errors)}")
    return records


def _load_orders(ws, sheets: dict[str, SheetLoad], log: list[str]) -> list[OrderRecord]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("订单 is empty")
    headers = [None if h is None else str(h).strip() for h in rows[0]]
    index = _header_index(headers)
    missing = [src for src, _, _ in ORDER_FIELDS if src not in index]
    if missing:
        raise KeyError(f"订单 missing columns: {missing}")
    records: list[OrderRecord] = []
    errors: list[dict[str, Any]] = []
    for excel_row, raw in enumerate(rows[1:], start=2):
        if _row_empty(raw):
            continue
        try:
            order_id = as_string_id(_cell(raw, index, "订单号"), field="order_id")
            session_id = as_string_id(_cell(raw, index, "会话ID"), field="session_id")
            if not order_id or not session_id:
                raise IdCoercionError(f"订单 row {excel_row}: missing primary key")
            placed_raw = _cell(raw, index, "下单时间")
            paid_raw = _cell(raw, index, "付款时间")
            shipped_raw = _cell(raw, index, "发货时间")
            placed = parse_shanghai(placed_raw, field="placed_time")
            if placed is None:
                raise TimeParseError(f"订单 row {excel_row}: 下单时间 is required")
            records.append(
                OrderRecord(
                    order_id=order_id,
                    session_id=session_id,
                    buyer_nick=as_string_id(_cell(raw, index, "买家昵称"), field="buyer_nick") or "",
                    shop=as_string(_cell(raw, index, "店铺")),
                    sku_code=as_string_id(_cell(raw, index, "商品货号"), field="sku_code"),
                    product_name=as_string(_cell(raw, index, "商品名称")),
                    quantity=as_int(_cell(raw, index, "数量"), field="quantity"),
                    unit_price=as_number(_cell(raw, index, "单价(元)"), field="unit_price"),
                    paid_amount=as_number(_cell(raw, index, "实付金额(元)"), field="paid_amount"),
                    order_status_raw=as_string(_cell(raw, index, "订单状态")),
                    placed_time=placed,
                    placed_time_raw=as_string(placed_raw) or "",
                    paid_time=parse_shanghai(paid_raw, field="paid_time"),
                    paid_time_raw=as_string(paid_raw),
                    paid_time_note=time_annotation(paid_raw),
                    shipped_time=parse_shanghai(shipped_raw, field="shipped_time"),
                    shipped_time_raw=as_string(shipped_raw),
                    courier=as_string(_cell(raw, index, "快递公司")),
                    tracking_no=as_string_id(_cell(raw, index, "物流单号"), field="tracking_no"),
                    receiver_province=as_string(_cell(raw, index, "收货省")),
                    receiver_city=as_string(_cell(raw, index, "收货市")),
                    gift=as_string(_cell(raw, index, "赠品")),
                    buyer_note=as_string(_cell(raw, index, "买家留言")),
                    source_sheet=ORDER_SHEET,
                    source_row=excel_row,
                )
            )
        except (IdCoercionError, TimeParseError, KeyError, ValueError) as exc:
            errors.append({"sheet": ORDER_SHEET, "row": excel_row, "error": str(exc)})
    sheets[ORDER_SHEET] = SheetLoad(
        name=ORDER_SHEET,
        headers=[h for h in headers if h],
        row_count=len(records),
        canonical_fields=[dst for _, dst, _ in ORDER_FIELDS],
        parse_errors=errors,
    )
    log.append(f"loaded {ORDER_SHEET} rows={len(records)} errors={len(errors)}")
    return records


def _load_workorders(ws, sheet_name: str, sheets: dict[str, SheetLoad], log: list[str]) -> list[WorkorderRecord]:
    meta = WORKORDER_SHEETS[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{sheet_name} is empty")
    headers = [None if h is None else str(h).strip() for h in rows[0]]
    index = _header_index(headers)
    status_field = meta["status_field"]
    needed = [src for src, _, _ in WORKORDER_COMMON_FIELDS] + [status_field] + [src for src, _, _ in meta["extra"]]
    missing = [name for name in needed if name not in index]
    if missing:
        raise KeyError(f"{sheet_name} missing columns: {missing}")
    records: list[WorkorderRecord] = []
    errors: list[dict[str, Any]] = []
    canonical_fields = [dst for _, dst, _ in WORKORDER_COMMON_FIELDS] + ["status_raw", "kind"] + [
        dst for _, dst, _ in meta["extra"]
    ]
    for excel_row, raw in enumerate(rows[1:], start=2):
        if _row_empty(raw):
            continue
        try:
            workorder_id = as_string_id(_cell(raw, index, "工单号"), field="workorder_id")
            session_id = as_string_id(_cell(raw, index, "会话ID"), field="session_id")
            if not workorder_id or not session_id:
                raise IdCoercionError(f"{sheet_name} row {excel_row}: missing primary key")
            create_raw = _cell(raw, index, "创建时间")
            create_time = parse_shanghai(create_raw, field="create_time")
            if create_time is None:
                raise TimeParseError(f"{sheet_name} row {excel_row}: 创建时间 is required")
            complete_raw = _cell(raw, index, "完成时间")
            extra = {}
            for src, dst, typ in meta["extra"]:
                extra[dst] = _convert(typ, _cell(raw, index, src), field=dst)
            records.append(
                WorkorderRecord(
                    workorder_id=workorder_id,
                    session_id=session_id,
                    order_id=as_string_id(_cell(raw, index, "关联订单号"), field="order_id"),
                    buyer_nick=as_string_id(_cell(raw, index, "买家昵称"), field="buyer_nick"),
                    shop=as_string(_cell(raw, index, "店铺")),
                    kind=meta["kind"],
                    source_sheet=sheet_name,
                    status_raw=as_string(_cell(raw, index, status_field)),
                    handler=as_string(_cell(raw, index, "处理人")),
                    create_time=create_time,
                    create_time_raw=as_string(create_raw) or "",
                    complete_time=parse_shanghai(complete_raw, field="complete_time"),
                    complete_time_raw=as_string(complete_raw),
                    extra=extra,
                    source_row=excel_row,
                )
            )
        except (IdCoercionError, TimeParseError, KeyError, ValueError) as exc:
            errors.append({"sheet": sheet_name, "row": excel_row, "error": str(exc)})
    sheets[sheet_name] = SheetLoad(
        name=sheet_name,
        headers=[h for h in headers if h],
        row_count=len(records),
        canonical_fields=canonical_fields,
        parse_errors=errors,
    )
    log.append(f"loaded {sheet_name} rows={len(records)} kind={meta['kind']} errors={len(errors)}")
    return records


def _row_empty(raw: tuple[Any, ...]) -> bool:
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in raw)
